"""
Embed remote brief images directly into an Excel workbook.

The Giro brief templates use Excel IMAGE() formulas pointing at Azure Blob
Storage. Those formulas often render as blank cells in Google Sheets because of
embedding/CORS restrictions. This script downloads each referenced image and
inserts it into the workbook so the pictures are visible in Excel and after
uploading to Google Sheets.

Usage:
    python embed_brief_images.py "path/to/brief.xlsx"
    python embed_brief_images.py "path/to/brief.xlsx" --output "path/to/out.xlsx"
"""
from __future__ import annotations

import argparse
import shutil
import sys
import tempfile
import urllib.error
import urllib.parse
import urllib.request
from pathlib import Path

import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import column_index_from_string, get_column_letter

IMAGE_SLOTS = (
    ("B", "A", "E"),
    ("G", "F", "H"),
    ("J", "I", "K"),
    ("M", "L", "N"),
)

MAX_IMAGE_WIDTH = 240
MAX_IMAGE_HEIGHT = 220


def build_image_url(base: str, folder: str, frame, file_format: str, token: str, digits: int) -> str:
    pad = "00000" if digits == 5 else "000000"
    frame_text = str(int(frame)).zfill(len(pad))
    base = base.rstrip("/")
    folder = str(folder).strip().strip("/")
    token = str(token).strip()
    if token and not token.startswith("?"):
        token = "?" + token

    encoded_folder = "/".join(
        urllib.parse.quote(part, safe="")
        for part in folder.split("/")
        if part
    )
    return f"{base}/{encoded_folder}/{frame_text}.{file_format}{token}"


def download_image(url: str, cache_dir: Path) -> Path | None:
    cache_name = str(abs(hash(url)))
    cached = cache_dir / f"{cache_name}.jpg"
    if cached.exists() and cached.stat().st_size > 0:
        return cached

    request = urllib.request.Request(url, headers={"User-Agent": "brief-image-embedder/1.0"})
    try:
        with urllib.request.urlopen(request, timeout=30) as response:
            data = response.read()
    except (urllib.error.URLError, TimeoutError, OSError):
        return None

    if not data:
        return None

    cached.write_bytes(data)
    return cached


def fit_image(image: XLImage) -> None:
    width = image.width or MAX_IMAGE_WIDTH
    height = image.height or MAX_IMAGE_HEIGHT
    scale = min(MAX_IMAGE_WIDTH / width, MAX_IMAGE_HEIGHT / height, 1.0)
    image.width = int(width * scale)
    image.height = int(height * scale)


def embed_images(workbook_path: Path, output_path: Path) -> tuple[int, int]:
    wb = openpyxl.load_workbook(workbook_path)
    if "Example images" not in wb.sheetnames or "Summary" not in wb.sheetnames:
        raise ValueError("Workbook must contain 'Summary' and 'Example images' sheets")

    summary = wb["Summary"]
    ws = wb["Example images"]

    base = summary["C2"].value
    token = summary["C3"].value
    file_format = summary["C4"].value
    digits = int(summary["C5"].value)

    if not base or not token or not file_format:
        raise ValueError("Summary sheet is missing Azure settings in C2:C5")

    cache_dir = Path(tempfile.mkdtemp(prefix="brief_images_"))
    embedded = 0
    failed = 0

    try:
        for frame_col, folder_col, image_col in IMAGE_SLOTS:
            frame_idx = column_index_from_string(frame_col)
            folder_idx = column_index_from_string(folder_col)

            for row in range(3, ws.max_row + 1):
                frame = ws.cell(row, frame_idx).value
                folder = ws.cell(row, folder_idx).value
                if frame in (None, ""):
                    continue

                cell_ref = f"{image_col}{row}"
                url = build_image_url(base, folder, frame, file_format, token, digits)
                image_path = download_image(url, cache_dir)
                if image_path is None:
                    ws[cell_ref].value = f"Image unavailable: {url}"
                    failed += 1
                    continue

                ws[cell_ref].value = None
                xl_image = XLImage(str(image_path))
                fit_image(xl_image)
                ws.add_image(xl_image, cell_ref)
                embedded += 1

        wb.save(output_path)
    finally:
        shutil.rmtree(cache_dir, ignore_errors=True)

    return embedded, failed


def main() -> None:
    parser = argparse.ArgumentParser(description="Embed Azure brief images into an Excel workbook")
    parser.add_argument("workbook", help="Path to the source .xlsx brief")
    parser.add_argument(
        "--output",
        help="Output .xlsx path (default: overwrite source file)",
    )
    args = parser.parse_args()

    workbook_path = Path(args.workbook).resolve()
    if not workbook_path.exists():
        print(f"Error: file not found: {workbook_path}")
        sys.exit(1)

    output_path = Path(args.output).resolve() if args.output else workbook_path
    if output_path != workbook_path:
        shutil.copy2(workbook_path, output_path)

    print(f"Embedding images into: {output_path}")
    embedded, failed = embed_images(workbook_path if output_path == workbook_path else output_path, output_path)
    print(f"Done. Embedded: {embedded}, failed: {failed}")


if __name__ == "__main__":
    main()
