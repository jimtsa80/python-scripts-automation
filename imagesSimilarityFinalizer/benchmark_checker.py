#!/usr/bin/env python3
"""
i-sieve Benchmark Checker
=========================
Compares new match annotation file(s) against the Benchmark Reference.

Usage:
    python benchmark_checker.py <match_file.xlsx> [part2.xlsx] [--ref Benchmark_Reference.xlsx]
    python benchmark_checker.py part1.xlsx part2.xlsx --roster tournament_roster.xlsx

Examples:
    python benchmark_checker.py match_part1.xlsx
    python benchmark_checker.py match_part1.xlsx match_part2.xlsx
    python benchmark_checker.py match_part1.xlsx --ref /path/to/Benchmark_Reference.xlsx
    python benchmark_checker.py part1.xlsx part2.xlsx --roster ICC_T20W_roster.xlsx --export
    python benchmark_checker.py part1.xlsx part2.xlsx --half-match   # ref = full T20, you annotate 50%
"""

import sys
import os
import argparse
import pandas as pd
import warnings
warnings.filterwarnings("ignore")

# ── ANSI colors for terminal output ──────────────────────────────────────────
RED    = "\033[91m"
YELLOW = "\033[93m"
GREEN  = "\033[92m"
CYAN   = "\033[96m"
MAGENTA = "\033[95m"
BOLD   = "\033[1m"
DIM    = "\033[2m"
RESET  = "\033[0m"

# ── Thresholds ────────────────────────────────────────────────────────────────
# Low vs scaled benchmark
WARN_BRAND_PCT    = 0.40   # <40%  → RED warning (significantly below)
CAUTION_BRAND_PCT = 0.70   # <70%  → YELLOW caution (below)

# High vs scaled benchmark (possible over-count / duplicate merge)
WARN_HIGH_BRAND_PCT = 2.0  # ≥200% → HIGH warning (significantly above)

# For asset-level (cross-brand) comparison, use percentile bands
# Below P25 → warning, P25-Median → caution, above Median → ok


def load_match_files(*paths: str) -> pd.DataFrame:
    """Load one or more match annotation xlsx files and combine them."""
    frames = []
    for p in paths:
        if not os.path.isfile(p):
            sys.exit(f"{RED}ERROR:{RESET} File not found: {p}")
        df = pd.read_excel(p, header=0)
        frames.append(df)
        print(f"  {DIM}Loaded:{RESET} {os.path.basename(p)}  ({len(df):,} rows)")
    combined = pd.concat(frames, ignore_index=True)
    # Rename columns defensively
    combined.columns = [c.strip() for c in combined.columns]
    required = {'Brand', 'Location', 'Duration'}
    missing = required - set(combined.columns)
    if missing:
        sys.exit(f"{RED}ERROR:{RESET} Missing columns in match file: {missing}")
    combined['Duration'] = pd.to_numeric(combined['Duration'], errors='coerce').fillna(0)
    return combined


def load_benchmark(ref_path: str):
    """Load Asset-level and Brand×Asset benchmark sheets."""
    if not os.path.isfile(ref_path):
        sys.exit(f"{RED}ERROR:{RESET} Benchmark reference not found: {ref_path}\n"
                 f"  Generate it first with the i-sieve benchmark builder.")

    bm_asset = pd.read_excel(ref_path, sheet_name="Asset Benchmarks", header=4)
    bm_asset.columns = [c.strip() for c in bm_asset.columns]
    # Rename the tricky spaced column names
    col_map = {}
    for c in bm_asset.columns:
        if 'P25' in c:   col_map[c] = 'P25'
        if 'P75' in c:   col_map[c] = 'P75'
        if 'Asset' in c: col_map[c] = 'Asset'
    bm_asset.rename(columns=col_map, inplace=True)

    bm_brand = pd.read_excel(ref_path, sheet_name="Brand x Asset Benchmarks", header=4)
    bm_brand.columns = [c.strip() for c in bm_brand.columns]
    col_map2 = {}
    for c in bm_brand.columns:
        if 'Asset' in c: col_map2[c] = 'Asset'
        if 'Benchmark' in c: col_map2[c] = 'Benchmark'
    bm_brand.rename(columns=col_map2, inplace=True)
    bm_brand['Brand']  = bm_brand['Brand'].str.strip()
    bm_brand['Asset']  = bm_brand['Asset'].str.strip()
    bm_brand['Benchmark'] = pd.to_numeric(bm_brand['Benchmark'], errors='coerce')

    return bm_asset, bm_brand


def _pick_column(columns, candidates):
    """Return first column name that matches any candidate (case-insensitive)."""
    lower = {c.lower().strip(): c for c in columns}
    for cand in candidates:
        if cand.lower() in lower:
            return lower[cand.lower()]
    return None


def load_tournament_roster(roster_path: str) -> pd.DataFrame:
    """
    Load tournament brand/asset roster (Brand, Asset, Example found).
    Example found = 1 → placement exists in this tournament and should appear in results.
    Example found = 0 → not expected for this tournament (benchmark warnings suppressed).
    """
    if not os.path.isfile(roster_path):
        sys.exit(f"{RED}ERROR:{RESET} Tournament roster not found: {roster_path}")

    df = pd.read_excel(roster_path, header=0)
    df.columns = [str(c).strip() for c in df.columns]

    brand_col = _pick_column(df.columns, ["Brand", "Brands"])
    asset_col = _pick_column(df.columns, ["Asset", "Assets", "Location"])
    found_col = _pick_column(df.columns, ["Example found", "Example Found", "Expected", "Found"])

    missing_cols = []
    if not brand_col:
        missing_cols.append("Brand(s)")
    if not asset_col:
        missing_cols.append("Asset(s)/Location")
    if not found_col:
        missing_cols.append("Example found")
    if missing_cols:
        sys.exit(f"{RED}ERROR:{RESET} Roster missing columns: {missing_cols}\n"
                 f"  Found: {list(df.columns)}")

    roster = pd.DataFrame({
        "Brand": df[brand_col].astype(str).str.strip(),
        "Asset": df[asset_col].astype(str).str.strip(),
        "Expected": pd.to_numeric(df[found_col], errors="coerce").fillna(0).astype(int),
    })
    roster = roster[(roster["Brand"] != "") & (roster["Asset"] != "")]
    roster = roster.drop_duplicates(subset=["Brand", "Asset"], keep="first")

    n_exp = int((roster["Expected"] == 1).sum())
    n_skip = int((roster["Expected"] == 0).sum())
    print(f"  {DIM}Roster:{RESET} {os.path.basename(roster_path)}  "
          f"({len(roster)} placements: {n_exp} expected, {n_skip} not expected)")
    return roster


def summarise_match(df: pd.DataFrame) -> pd.DataFrame:
    """Aggregate total seconds per Brand × Location."""
    summary = (
        df.groupby(['Brand', 'Location'])['Duration']
        .sum()
        .reset_index()
    )
    summary.columns = ['Brand', 'Asset', 'Total_Sec']
    summary['Brand'] = summary['Brand'].str.strip()
    summary['Asset'] = summary['Asset'].str.strip()
    return summary


def _scale_ref(value, fraction: float):
    """Scale a benchmark value for partial-match coverage (e.g. 0.5 = half T20)."""
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return value
    return float(value) * fraction


def _ref_label(full_ref, adj_ref, fraction: float) -> str:
    """Format benchmark reference for messages."""
    if fraction >= 0.999:
        return f"{adj_ref:.0f}s"
    return f"{adj_ref:.0f}s ({full_ref:.0f}s full × {fraction*100:.0f}%)"


def _by_match_sec_desc(items: list) -> list:
    """Sort result rows by This Match seconds, largest first."""
    return sorted(items, key=lambda x: x.get("match_sec", 0), reverse=True)


def band_label(value, p25, median, p75):
    """Return performance band given percentile thresholds."""
    if pd.isna(p25) or pd.isna(median) or pd.isna(p75):
        return None
    if value < p25:
        return 'BELOW_P25'
    elif value < median:
        return 'P25_TO_MEDIAN'
    elif value < p75:
        return 'MEDIAN_TO_P75'
    else:
        return 'ABOVE_P75'


def run_roster_checks(summary: pd.DataFrame, roster: pd.DataFrame):
    """
    Cross-check match totals against the tournament roster.
    Returns dicts/lists used by the report and export.
    """
    match_lookup = {
        (r["Brand"], r["Asset"]): float(r["Total_Sec"])
        for _, r in summary.iterrows()
    }
    roster_lookup = {
        (r["Brand"], r["Asset"]): int(r["Expected"])
        for _, r in roster.iterrows()
    }

    missing_expected = []      # roster Expected=1 but 0 sec in match
    not_expected_found = []    # roster Expected=0 but >0 sec in match
    not_in_roster = []         # match has sec>0 but pair absent from roster
    roster_ok = []             # Expected=1 and sec>0

    for (brand, asset), expected in roster_lookup.items():
        sec = match_lookup.get((brand, asset), 0.0)
        if expected == 1:
            if sec <= 0:
                missing_expected.append({
                    "brand": brand, "asset": asset, "match_sec": sec,
                    "msg": f"Expected in tournament but NOT FOUND in match (0s)"
                })
            else:
                roster_ok.append({
                    "brand": brand, "asset": asset, "match_sec": sec,
                    "msg": f"{sec:.0f}s  ✓  expected & detected"
                })
        elif sec > 0:
            not_expected_found.append({
                "brand": brand, "asset": asset, "match_sec": sec,
                "msg": f"{sec:.0f}s  →  roster says not expected (Example found=0) but appears in match"
            })

    for (brand, asset), sec in match_lookup.items():
        if sec > 0 and (brand, asset) not in roster_lookup:
            not_in_roster.append({
                "brand": brand, "asset": asset, "match_sec": sec,
                "msg": f"{sec:.0f}s  →  not listed in tournament roster"
            })

    missing_expected = _by_match_sec_desc(missing_expected)
    not_expected_found = _by_match_sec_desc(not_expected_found)
    not_in_roster = _by_match_sec_desc(not_in_roster)
    roster_ok = _by_match_sec_desc(roster_ok)

    sep = "─" * 72
    print(f"\n{BOLD}{'═'*72}{RESET}")
    print(f"{BOLD}  TOURNAMENT ROSTER CHECK{RESET}")
    print(f"  Roster placements: {len(roster_lookup)}   "
          f"Expected: {sum(1 for v in roster_lookup.values() if v == 1)}")
    print(f"{BOLD}{'═'*72}{RESET}\n")

    if missing_expected:
        print(f"{RED}{BOLD}⛔  MISSING (expected in tournament, not in match)  ({len(missing_expected)}){RESET}")
        print(f"{DIM}{sep}{RESET}")
        for m in missing_expected:
            print(f"  {RED}▶ {m['brand']} — {m['asset']}{RESET}")
            print(f"     {m['msg']}")
        print()
    else:
        print(f"{GREEN}✓  All roster-expected placements found in match{RESET}\n")

    if not_in_roster:
        print(f"{YELLOW}{BOLD}⚠   NOT IN ROSTER (in match, absent from tournament list)  ({len(not_in_roster)}){RESET}")
        print(f"{DIM}{sep}{RESET}")
        for n in not_in_roster:
            print(f"  {YELLOW}▷ {n['brand']} — {n['asset']}{RESET}")
            print(f"     {n['msg']}")
        print()

    if not_expected_found:
        print(f"{CYAN}{BOLD}ℹ   NOT EXPECTED (roster=0 but detected)  ({len(not_expected_found)}){RESET}")
        print(f"{DIM}{sep}{RESET}")
        for n in not_expected_found:
            print(f"  {CYAN}· {n['brand']} — {n['asset']}{RESET}")
            print(f"     {n['msg']}")
        print()

    print(f"{DIM}{sep}{RESET}")
    print(f"  {BOLD}ROSTER SUMMARY:{RESET}  "
          f"{RED}⛔ {len(missing_expected)} missing{RESET}  "
          f"{GREEN}✅ {len(roster_ok)} expected & found{RESET}  "
          f"{YELLOW}⚠ {len(not_in_roster)} not in roster{RESET}  "
          f"{CYAN}ℹ {len(not_expected_found)} unexpected{RESET}")
    print(f"{BOLD}{'═'*72}{RESET}\n")

    return {
        "missing_expected": missing_expected,
        "not_expected_found": not_expected_found,
        "not_in_roster": not_in_roster,
        "roster_ok": roster_ok,
        "roster_lookup": roster_lookup,
    }


def run_checks(summary: pd.DataFrame, bm_asset: pd.DataFrame, bm_brand: pd.DataFrame,
               match_name: str, roster_lookup: dict | None = None,
               coverage_fraction: float = 1.0):
    """Run all benchmark comparisons and print a structured report."""

    # Build lookup dicts
    brand_lookup = {}   # (Brand, Asset) → benchmark_sec
    for _, r in bm_brand.iterrows():
        if pd.notna(r['Benchmark']):
            brand_lookup[(r['Brand'], r['Asset'])] = r['Benchmark']

    asset_lookup = {}   # Asset → {P25, Median, P75, Mean, N}
    for _, r in bm_asset.iterrows():
        asset_lookup[r['Asset']] = {
            'N':      r.get('N Brands', None),
            'Min':    r.get('Min', None),
            'P25':    r.get('P25', None),
            'Median': r.get('Median', None),
            'Mean':   r.get('Mean', None),
            'P75':    r.get('P75', None),
            'Max':    r.get('Max', None),
        }

    warnings_list  = []
    cautions_list  = []
    high_list      = []  # significantly above benchmark
    ok_list        = []
    no_ref_list    = []

    skipped_roster = []  # Expected=0 in roster — benchmark skipped

    for _, row in summary.iterrows():
        brand = row['Brand']
        asset = row['Asset']
        total_sec = row['Total_Sec']

        if roster_lookup is not None:
            expected = roster_lookup.get((brand, asset))
            if expected == 0:
                skipped_roster.append({
                    "brand": brand, "asset": asset, "match_sec": total_sec,
                    "msg": f"{total_sec:.0f}s  →  skipped benchmark (not expected in tournament roster)"
                })
                continue
            if expected is None and total_sec > 0:
                # handled in roster section; still run benchmark if data exists
                pass

        bm_full  = brand_lookup.get((brand, asset))   # brand-specific benchmark (full match)
        asset_bm = asset_lookup.get(asset)            # cross-brand percentiles (full match)

        # ── 1. Brand-specific comparison ──────────────────────────────────
        if bm_full is not None:
            bm_val = _scale_ref(bm_full, coverage_fraction)
            ref_str = _ref_label(bm_full, bm_val, coverage_fraction)
            pct = total_sec / bm_val if bm_val > 0 else 0
            if pct >= WARN_HIGH_BRAND_PCT:
                high_list.append({
                    'level': 'HIGH', 'brand': brand, 'asset': asset,
                    'match_sec': total_sec, 'ref': bm_val, 'ref_full': bm_full,
                    'msg': f"{total_sec:.0f}s vs benchmark {ref_str}  →  {pct*100:.0f}% of expected  ↑  SIGNIFICANTLY ABOVE"
                })
            elif pct < WARN_BRAND_PCT:
                warnings_list.append({
                    'level': 'WARNING', 'brand': brand, 'asset': asset,
                    'match_sec': total_sec, 'ref': bm_val, 'ref_full': bm_full,
                    'msg': f"{total_sec:.0f}s vs benchmark {ref_str}  →  {pct*100:.0f}% of expected  ⚠  SIGNIFICANTLY BELOW"
                })
            elif pct < CAUTION_BRAND_PCT:
                cautions_list.append({
                    'level': 'CAUTION', 'brand': brand, 'asset': asset,
                    'match_sec': total_sec, 'ref': bm_val, 'ref_full': bm_full,
                    'msg': f"{total_sec:.0f}s vs benchmark {ref_str}  →  {pct*100:.0f}% of expected  ↓  below benchmark"
                })
            else:
                ok_list.append({
                    'level': 'OK', 'brand': brand, 'asset': asset,
                    'match_sec': total_sec, 'ref': bm_val, 'ref_full': bm_full,
                    'msg': f"{total_sec:.0f}s vs benchmark {ref_str}  →  {pct*100:.0f}%  ✓"
                })

        # ── 2. Asset-level percentile band (cross-brand) ───────────────────
        elif asset_bm is not None:
            p25 = _scale_ref(asset_bm['P25'], coverage_fraction)
            med = _scale_ref(asset_bm['Median'], coverage_fraction)
            p75 = _scale_ref(asset_bm['P75'], coverage_fraction)
            n   = asset_bm['N']
            band = band_label(total_sec, p25, med, p75)
            if coverage_fraction < 0.999:
                band_str = (f"P25={p25:.0f}  Median={med:.0f}  P75={p75:.0f}  "
                            f"(N={n:.0f} brands, scaled ×{coverage_fraction:.0%})")
            else:
                band_str = f"P25={p25:.0f}  Median={med:.0f}  P75={p75:.0f}  (N={n:.0f} brands)"
            if band == 'BELOW_P25':
                warnings_list.append({
                    'level': 'WARNING', 'brand': brand, 'asset': asset,
                    'match_sec': total_sec, 'ref': None,
                    'msg': f"{total_sec:.0f}s  →  BELOW P25  [{band_str}]  ⚠  no brand-specific ref, using cross-brand"
                })
            elif band == 'P25_TO_MEDIAN':
                cautions_list.append({
                    'level': 'CAUTION', 'brand': brand, 'asset': asset,
                    'match_sec': total_sec, 'ref': None,
                    'msg': f"{total_sec:.0f}s  →  P25→Median  [{band_str}]  ↓  lower half"
                })
            elif band in ('MEDIAN_TO_P75', 'ABOVE_P75'):
                ref_for_high = med if med and med > 0 else p75
                if ref_for_high and ref_for_high > 0 and total_sec / ref_for_high >= WARN_HIGH_BRAND_PCT:
                    high_list.append({
                        'level': 'HIGH', 'brand': brand, 'asset': asset,
                        'match_sec': total_sec, 'ref': ref_for_high,
                        'msg': (f"{total_sec:.0f}s  →  {total_sec/ref_for_high*100:.0f}% of median "
                                f"[{band_str}]  ↑  SIGNIFICANTLY ABOVE (cross-brand)")
                    })
                else:
                    ok_list.append({
                        'level': 'OK', 'brand': brand, 'asset': asset,
                        'match_sec': total_sec, 'ref': None,
                        'msg': f"{total_sec:.0f}s  →  {band.replace('_',' ')}  [{band_str}]  ✓"
                    })
            else:
                no_ref_list.append({'brand': brand, 'asset': asset, 'match_sec': total_sec})
        else:
            no_ref_list.append({'brand': brand, 'asset': asset, 'match_sec': total_sec})

    # ── PRINT REPORT ──────────────────────────────────────────────────────────
    sep = "─" * 72

    print(f"\n{BOLD}{'═'*72}{RESET}")
    print(f"{BOLD}  i-sieve BENCHMARK CHECK REPORT{RESET}")
    print(f"  Match:  {CYAN}{match_name}{RESET}")
    print(f"  Brands: {len(summary['Brand'].unique())}   Placements tracked: {len(summary)}")
    if coverage_fraction < 0.999:
        print(f"  Coverage: {CYAN}{coverage_fraction*100:.0f}%{RESET} of full match  "
              f"{DIM}(benchmark reference scaled from full T20){RESET}")
    print(f"{BOLD}{'═'*72}{RESET}\n")

    warnings_list = _by_match_sec_desc(warnings_list)
    cautions_list = _by_match_sec_desc(cautions_list)
    high_list = _by_match_sec_desc(high_list)
    ok_list = _by_match_sec_desc(ok_list)
    no_ref_list = _by_match_sec_desc(no_ref_list)
    skipped_roster = _by_match_sec_desc(skipped_roster)

    # ── WARNINGS ──
    if warnings_list:
        print(f"{RED}{BOLD}⛔  LOW (significantly below benchmark)  ({len(warnings_list)}){RESET}")
        print(f"{DIM}{sep}{RESET}")
        for w in warnings_list:
            print(f"  {RED}▶ {w['brand']} — {w['asset']}{RESET}")
            print(f"     {w['msg']}")
        print()
    else:
        print(f"{GREEN}✓  No critical warnings{RESET}\n")

    # ── CAUTIONS ──
    if cautions_list:
        print(f"{YELLOW}{BOLD}⚠   CAUTIONS  ({len(cautions_list)}){RESET}")
        print(f"{DIM}{sep}{RESET}")
        for c in cautions_list:
            print(f"  {YELLOW}▷ {c['brand']} — {c['asset']}{RESET}")
            print(f"     {c['msg']}")
        print()

    # ── HIGH (above benchmark) ──
    if high_list:
        print(f"{MAGENTA}{BOLD}📈  HIGH (significantly above benchmark)  ({len(high_list)}){RESET}")
        print(f"{DIM}{sep}{RESET}")
        for h in high_list:
            print(f"  {MAGENTA}▲ {h['brand']} — {h['asset']}{RESET}")
            print(f"     {h['msg']}")
        print()

    # ── OK ──
    print(f"{GREEN}{BOLD}✅  WITHIN BENCHMARK  ({len(ok_list)}){RESET}")
    print(f"{DIM}{sep}{RESET}")
    for o in ok_list:
        print(f"  {GREEN}· {o['brand']} — {o['asset']}{RESET}")
        print(f"     {o['msg']}")
    print()

    # ── NEW / NO REFERENCE ──
    if no_ref_list:
        print(f"{CYAN}{BOLD}ℹ   NEW PLACEMENTS (no benchmark reference)  ({len(no_ref_list)}){RESET}")
        print(f"{DIM}{sep}{RESET}")
        for n in no_ref_list:
            print(f"  {CYAN}· {n['brand']} — {n['asset']}{RESET}  ({n['match_sec']:.0f}s)")
        print()

    if skipped_roster:
        print(f"{DIM}{BOLD}⊘   BENCHMARK SKIPPED (roster Example found=0)  ({len(skipped_roster)}){RESET}")
        print(f"{DIM}{sep}{RESET}")
        for s in skipped_roster:
            print(f"  {DIM}· {s['brand']} — {s['asset']}{RESET}")
            print(f"     {s['msg']}")
        print()

    # ── SUMMARY COUNTS ──
    total = (len(warnings_list) + len(cautions_list) + len(high_list)
             + len(ok_list) + len(no_ref_list))
    print(f"{DIM}{sep}{RESET}")
    print(f"  {BOLD}BENCHMARK SUMMARY:{RESET}  "
          f"{RED}⛔ {len(warnings_list)} low{RESET}  "
          f"{YELLOW}⚠ {len(cautions_list)} cautions{RESET}  "
          f"{MAGENTA}📈 {len(high_list)} high{RESET}  "
          f"{GREEN}✅ {len(ok_list)} ok{RESET}  "
          f"{CYAN}ℹ {len(no_ref_list)} new{RESET}  "
          f"{DIM}⊘ {len(skipped_roster)} roster-skipped{RESET}  "
          f"  (total {total} placements)")
    print(f"{BOLD}{'═'*72}{RESET}\n")

    return warnings_list, cautions_list, high_list, ok_list, no_ref_list, skipped_roster


def export_roster_sheet(wb, roster: pd.DataFrame, summary: pd.DataFrame, roster_results: dict):
    """Add Tournament Roster sheet with match seconds and status per row."""
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    HDR = "1A2332"
    WHITE = "FFFFFF"
    R_BG = "FFF0F0"
    G_BG = "F0FFF4"
    Y_BG = "FFFBF0"
    C_BG = "F0FBFF"
    DIM_BG = "F6F8FA"

    def fill(c):
        return PatternFill("solid", start_color=c, end_color=c)

    def thin():
        s = Side(style='thin', color='D0D7DE')
        return Border(left=s, right=s, top=s, bottom=s)

    match_lookup = {
        (r["Brand"], r["Asset"]): float(r["Total_Sec"])
        for _, r in summary.iterrows()
    }

    ws = wb.create_sheet("Tournament Roster")
    headers = ["Brand", "Asset", "Example found", "Match (sec)", "Status"]
    widths = [22, 40, 14, 14, 36]
    for i, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = Font(name='Arial', bold=True, color=WHITE, size=10)
        c.fill = fill(HDR)
        c.alignment = Alignment(horizontal='center', vertical='center')
        c.border = thin()
        ws.column_dimensions[get_column_letter(i)].width = w

    for ri, (_, row) in enumerate(roster.iterrows(), 2):
        brand, asset, expected = row["Brand"], row["Asset"], int(row["Expected"])
        sec = match_lookup.get((brand, asset), 0.0)

        if expected == 1 and sec <= 0:
            status, bg = "MISSING", R_BG
        elif expected == 0 and sec > 0:
            status, bg = "NOT EXPECTED (found)", C_BG
        elif expected == 1 and sec > 0:
            status, bg = "OK", G_BG
        else:
            status, bg = "Not expected", DIM_BG

        vals = [brand, asset, expected, sec if sec > 0 else "—", status]
        for ci, v in enumerate(vals, 1):
            c = ws.cell(row=ri, column=ci, value=v)
            c.fill = fill(bg)
            c.border = thin()
            c.alignment = Alignment(
                horizontal='center' if ci != 2 else 'left',
                vertical='center', wrap_text=(ci == 2)
            )
            if ci == 3 and v == 1:
                c.fill = fill(G_BG)
            elif ci == 3 and v == 0:
                c.fill = fill(R_BG)
        ws.row_dimensions[ri].height = 18

    ws.freeze_panes = "A2"


def export_report(summary, warnings_list, cautions_list, high_list, ok_list, no_ref_list,
                  match_name: str, out_path: str,
                  roster: pd.DataFrame | None = None, roster_results: dict | None = None):
    """Optionally export results to Excel."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    DARK   = "0D1117"
    TEAL   = "00C896"
    WHITE  = "FFFFFF"
    R_BG   = "FFF0F0"
    Y_BG   = "FFFBF0"
    G_BG   = "F0FFF4"
    C_BG   = "F0FBFF"
    P_BG   = "F3E8FF"
    HDR    = "1A2332"

    def fill(c): return PatternFill("solid", start_color=c, end_color=c)
    def font(bold=False, color="1A1A2E", size=10):
        return Font(name='Arial', bold=bold, color=color, size=size)
    def thin():
        s = Side(style='thin', color='D0D7DE')
        return Border(left=s, right=s, top=s, bottom=s)

    wb = Workbook()
    ws = wb.active
    ws.title = "Benchmark Check"

    # Title
    ws.merge_cells("A1:F1")
    ws["A1"] = f"i-sieve Benchmark Check — {match_name}"
    ws["A1"].font = Font(name='Arial', bold=True, color=TEAL, size=13)
    ws["A1"].fill = fill(DARK)
    ws["A1"].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 28

    # Headers
    headers = ["Status", "Brand", "Asset (Placement)", "This Match (sec)", "Reference (sec)", "Notes"]
    widths  = [14, 20, 38, 18, 18, 52]
    for i, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(row=2, column=i, value=h)
        c.font = Font(name='Arial', bold=True, color=WHITE, size=10)
        c.fill = fill(HDR)
        c.alignment = Alignment(horizontal='center', vertical='center')
        c.border = thin()
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[2].height = 22

    all_rows = (
        [("⛔ LOW",      R_BG, w) for w in _by_match_sec_desc(warnings_list)] +
        [("⚠ CAUTION",  Y_BG, w) for w in _by_match_sec_desc(cautions_list)] +
        [("📈 HIGH",     P_BG, w) for w in _by_match_sec_desc(high_list)] +
        [("✅ OK",        G_BG, w) for w in _by_match_sec_desc(ok_list)] +
        [("ℹ NEW",       C_BG, w) for w in _by_match_sec_desc(no_ref_list)]
    )

    for ri, (status, bg, row) in enumerate(all_rows, 3):
        ref_val = row.get('ref', None)
        msg     = row.get('msg', f"{row['match_sec']:.0f}s — no benchmark reference")
        vals = [status, row['brand'], row['asset'], row['match_sec'],
                ref_val if ref_val else "—", msg]
        for ci, v in enumerate(vals, 1):
            c = ws.cell(row=ri, column=ci, value=v)
            c.fill = fill(bg)
            c.border = thin()
            c.alignment = Alignment(horizontal='center' if ci != 6 else 'left',
                                    vertical='center', wrap_text=(ci==6))
            if ci == 1:
                c.font = Font(name='Arial', bold=True, size=10,
                              color=("C0392B" if "LOW" in status
                                     else "B7770D" if "CAUTION" in status
                                     else "6B21A8" if "HIGH" in status
                                     else "1A6B3C" if "OK" in status else "0D6E8F"))
            else:
                c.font = font()
            if ci in (4, 5) and isinstance(v, (int, float)):
                c.number_format = '#,##0.0'
        ws.row_dimensions[ri].height = 18

    ws.freeze_panes = "A3"

    if roster is not None and roster_results is not None:
        export_roster_sheet(wb, roster, summary, roster_results)

    wb.save(out_path)
    print(f"  {GREEN}Excel report saved:{RESET} {out_path}")


def main():
    parser = argparse.ArgumentParser(
        description="i-sieve Benchmark Checker — compare match annotations vs reference benchmarks"
    )
    parser.add_argument("files", nargs="+",
                        help="Match annotation xlsx file(s). Pass 2 for Part1 + Part2.")
    parser.add_argument("--ref", default="Benchmark_Reference.xlsx",
                        help="Path to Benchmark_Reference.xlsx (default: ./Benchmark_Reference.xlsx)")
    parser.add_argument("--roster", default=None,
                        help="Tournament roster xlsx (columns: Brand, Asset, Example found)")
    parser.add_argument("--coverage", type=float, default=1.0,
                        help="Fraction of full match annotated (default: 1.0). Use 0.5 for half T20.")
    parser.add_argument("--half-match", action="store_true",
                        help="Shorthand for --coverage 0.5 (reference is full match, you annotate half)")
    parser.add_argument("--export", action="store_true",
                        help="Also export results to an Excel report")
    parser.add_argument("--out", default=None,
                        help="Output Excel path (default: auto-named next to input)")
    args = parser.parse_args()

    coverage = 0.5 if args.half_match else args.coverage
    if not 0 < coverage <= 1.0:
        sys.exit(f"{RED}ERROR:{RESET} --coverage must be between 0 and 1 (got {coverage})")

    print(f"\n{BOLD}i-sieve Benchmark Checker{RESET}")
    if coverage < 0.999:
        print(f"  {DIM}Match coverage:{RESET} {coverage*100:.0f}% of full T20  "
              f"{DIM}(benchmarks scaled accordingly){RESET}")
    print(f"{DIM}Loading files...{RESET}")

    df     = load_match_files(*args.files)
    bm_asset, bm_brand = load_benchmark(args.ref)
    summary = summarise_match(df)

    roster = None
    roster_results = None
    roster_lookup = None
    if args.roster:
        roster = load_tournament_roster(args.roster)
        roster_lookup = {
            (r["Brand"], r["Asset"]): int(r["Expected"])
            for _, r in roster.iterrows()
        }
        roster_results = run_roster_checks(summary, roster)

    match_name = os.path.basename(args.files[0]).replace(".xlsx", "")
    if len(args.files) > 1:
        match_name += f" (+{len(args.files)-1} parts)"

    warnings_list, cautions_list, high_list, ok_list, no_ref_list, _skipped = run_checks(
        summary, bm_asset, bm_brand, match_name,
        roster_lookup=roster_lookup, coverage_fraction=coverage
    )

    if args.export:
        if args.out:
            out_path = args.out
        else:
            stem = os.path.splitext(os.path.basename(args.files[0]))[0]
            out_dir = os.path.dirname(os.path.abspath(args.files[0]))
            out_path = os.path.join(out_dir, f"{stem}_Check_Old.xlsx")
        export_report(summary, warnings_list, cautions_list, high_list, ok_list, no_ref_list,
                      match_name, out_path, roster=roster, roster_results=roster_results)


if __name__ == "__main__":
    main()
