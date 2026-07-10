import pandas as pd
import sys
import re
from datetime import datetime, timedelta
from openpyxl import load_workbook
from openpyxl.styles import NamedStyle

def extract_start_time(file_path):
    """
    Extracts a 4-digit time (HHMM) from the end of the filename.
    Returns a datetime object (with default date 1900-01-01) or defaults to 00:00.
    """
    print("Extracting start time from filename...")
    filename = file_path.split("_")[-1].replace(".xlsx", "")
    match = re.search(r'(\d{4})$', filename)
    if match:
        base_time = match.group(1)
        print(f"Found time in filename: {base_time}")
        try:
            return datetime.strptime(base_time, "%H%M")
        except ValueError:
            print("Failed to parse extracted time, defaulting to 00:00.")
    else:
        print("No valid time found in filename, defaulting to 00:00.")
    return datetime.strptime("0000", "%H%M")

def extract_frame_number(frame_str):
    """
    Extracts the numeric portion from the frame number.
    """
    print(f"Extracting frame number from: {frame_str}")
    frame_str = str(frame_str).strip()
    if frame_str.isdigit():
        print(f"Frame number is purely numeric: {frame_str}")
        return int(frame_str)
    
    match = re.search(r'_(\d+)$', frame_str)
    if match:
        print(f"Extracted numeric frame: {match.group(1)}")
        return int(match.group(1))
    
    print("No numeric frame found, defaulting to 0.")
    return 0

def process_excel(file_path):
    print(f"Processing file: {file_path}")
    print("Loading Excel file...")
    df = pd.read_excel(file_path)
    print("Initial data loaded:")
    print(df.head())

    # Extract start time from file name
    start_time = extract_start_time(file_path)
    print(f"Extracted start time: {start_time.strftime('%H:%M:%S')}")

    # Calculate "Time the brand is at screen"
    print("Calculating 'Time the brand is at screen'...")
    df['Time the brand is at screen'] = df['Sequence Frame Number'].apply(
        lambda x: start_time + timedelta(seconds=extract_frame_number(x))
    )
    df['Time the brand is at screen'] = pd.to_datetime(df['Time the brand is at screen']).dt.time
    print("Sample 'Time the brand is at screen' values:")
    print(df[['Sequence Frame Number', 'Time the brand is at screen']].head())

    # Group the DataFrame by relevant columns
    print("Grouping and aggregating data...")
    grouped = df.groupby(
        ['Brand', 'Location', 'Time the brand is at screen', 'Screen Location', 'Screen Size %', 'Sequence Frame Number'],
        as_index=False
    ).agg({
        'Duration': 'sum',
        'Total Hits': 'sum'
    })

    # Format Screen Size %
    grouped['Screen Size %'] = grouped['Screen Size %'].round(2)

    # Ensure Total Hits is at least equal to Duration
    grouped['Total Hits'] = grouped.apply(
        lambda row: max(row['Total Hits'], row['Duration']),
        axis=1
    )
    grouped['Average Hits'] = (grouped['Total Hits'] / grouped['Duration']).round(2)
    print("Grouped data sample:")
    print(grouped.head())

    # Sorting
    print("Sorting data by 'Sequence Frame Number'...")
    result_df = grouped.sort_values(by='Sequence Frame Number')

    # Reorder columns
    columns = result_df.columns.tolist()
    duration_index = columns.index('Duration')
    columns.insert(3, columns.pop(duration_index))
    columns.remove('Sequence Frame Number')
    columns.append('Sequence Frame Number')
    result_df = result_df[columns]
    print("Final column order:", columns)

    # Save the result
    print("Saving processed data back to Excel...")
    with pd.ExcelWriter(file_path, engine='openpyxl', mode='w') as writer:
        result_df.to_excel(writer, index=False)

    # Apply time formatting using openpyxl
    print("Applying time formatting in Excel...")
    wb = load_workbook(file_path)
    ws = wb.active
    time_style = NamedStyle(name="time_style", number_format="HH:MM:SS")
    time_col_index = result_df.columns.get_loc("Time the brand is at screen") + 1
    for row in ws.iter_rows(min_row=2, min_col=time_col_index, max_col=time_col_index, max_row=ws.max_row):
        for cell in row:
            cell.style = time_style
    wb.save(file_path)

    print("Process completed successfully.")
    print(f"Total Duration: {result_df['Duration'].sum()}")

if __name__ == "__main__":
    if len(sys.argv) != 2:
        print("Usage: python script.py <excel_file_path>")
        sys.exit(1)
    
    excel_file_path = sys.argv[1]
    process_excel(excel_file_path)