#!/usr/bin/env python3
"""
i-sieve Median Checker (current tournament pool)
================================================
Compare a new match file against annotations already in results/:

  (A) ALL VENUES — median seconds per Brand×Asset across every xlsx in results/
  (B) SAME VENUE — median/mean at the given venue folder (absolute sec comparison)

Use benchmark_checker.py for historical Benchmark_Reference.xlsx comparisons.
Use this script for comparisons against matches you have already finalised this event.

Usage:
    python median_checker.py match.xlsx --venue Edgbaston
    python median_checker.py part1.xlsx part2.xlsx --venue "Old Trafford" --export
    python median_checker.py match.xlsx --venue Edgbaston --results-dir results --export

Folder layout expected:
    results/
      Edgbaston/
        final_....xlsx
      Old Trafford/
        final_....xlsx
"""

from __future__ import annotations

import argparse
import os
import sys
from pathlib import Path

import pandas as pd
import warnings

warnings.filterwarnings("ignore")

# Reuse colours, thresholds, loaders from benchmark_checker
from benchmark_checker import (
    BOLD, CYAN, DIM, GREEN, MAGENTA, RED, RESET, YELLOW,
    CAUTION_BRAND_PCT,
    WARN_BRAND_PCT,
    WARN_HIGH_BRAND_PCT,
    _by_match_sec_desc,
    load_match_files,
    summarise_match,
)

SCRIPT_DIR = Path(__file__).resolve().parent
DEFAULT_RESULTS_DIR = SCRIPT_DIR / "results"


def _resolve_paths(paths: list[str]) -> list[Path]:
    return [Path(p).resolve() for p in paths]


def discover_pool_xlsx(results_dir: Path, venue: str | None, exclude: set[Path]) -> tuple[list[Path], list[Path]]:
    """
    Return (all_venue_files, same_venue_files).
    Excludes paths in `exclude` (the match being checked).
    """
    if not results_dir.is_dir():
        sys.exit(f"{RED}ERROR:{RESET} Results directory not found: {results_dir}")

    all_files: list[Path] = []
    for p in sorted(results_dir.rglob("*.xlsx")):
        if p.name.startswith("~$"):
            continue
        rp = p.resolve()
        if rp in exclude:
            continue
        all_files.append(rp)

    venue_files: list[Path] = []
    if venue:
        venue_dir = results_dir / venue
        if venue_dir.is_dir():
            for p in sorted(venue_dir.glob("*.xlsx")):
                if p.name.startswith("~$"):
                    continue
                rp = p.resolve()
                if rp not in exclude:
                    venue_files.append(rp)
        else:
            print(f"{YELLOW}Warning:{RESET} Venue folder not found: {venue_dir}")

    return all_files, venue_files


def load_summaries_from_files(files: list[Path]) -> tuple[pd.DataFrame, list[str]]:
    """Load each xlsx → Brand×Asset totals. Returns long dataframe + source names."""
    if not files:
        return pd.DataFrame(columns=["Brand", "Asset", "Total_Sec", "_source"]), []

    frames = []
    sources = []
    for fp in files:
        try:
            df = pd.read_excel(fp, header=0)
        except Exception as e:
            print(f"{YELLOW}  Skip:{RESET} {fp.name}  ({e})")
            continue
        df.columns = [str(c).strip() for c in df.columns]
        if not {"Brand", "Location", "Duration"}.issubset(df.columns):
            print(f"{YELLOW}  Skip:{RESET} {fp.name}  (missing Brand/Location/Duration)")
            continue
        df["Duration"] = pd.to_numeric(df["Duration"], errors="coerce").fillna(0)
        summary = summarise_match(df)
        summary["_source"] = fp.name
        frames.append(summary)
        sources.append(fp.name)
        print(f"  {DIM}Pool:{RESET} {fp.parent.name}/{fp.name}  ({len(summary)} placements)")

    if not frames:
        return pd.DataFrame(columns=["Brand", "Asset", "Total_Sec", "_source"]), []

    return pd.concat(frames, ignore_index=True), sources


def build_median_stats(long_df: pd.DataFrame) -> pd.DataFrame:
    """Per Brand×Asset: median, mean, count, min, max across match files."""
    if long_df.empty:
        return pd.DataFrame(columns=["Brand", "Asset", "median", "mean", "count", "min", "max"])

    stats = (
        long_df.groupby(["Brand", "Asset"], as_index=False)["Total_Sec"]
        .agg(median="median", mean="mean", count="count", min="min", max="max")
    )
    for col in ("median", "mean", "min", "max"):
        stats[col] = stats[col].round(1)
    return stats


def _classify_row(brand: str, asset: str, match_sec: float, ref: float, ref_label: str,
                  n_matches: int, context: str) -> tuple[str, dict]:
    """Return (bucket, row_dict) where bucket is low|caution|high|ok|no_ref."""
    base = {"brand": brand, "asset": asset, "match_sec": match_sec, "ref": ref,
            "n": n_matches, "context": context}

    if ref is None or pd.isna(ref) or ref <= 0 or n_matches == 0:
        return "no_ref", {**base, "msg": f"{match_sec:.0f}s  →  no reference in {context} pool"}

    pct = match_sec / ref
    ref_str = f"{ref:.0f}s (n={int(n_matches)})"

    if pct >= WARN_HIGH_BRAND_PCT:
        return "high", {
            **base,
            "msg": (f"{match_sec:.0f}s vs {ref_label} {ref_str}  →  {pct*100:.0f}%  "
                    f"↑  SIGNIFICANTLY ABOVE [{context}]")
        }
    if pct < WARN_BRAND_PCT:
        return "low", {
            **base,
            "msg": (f"{match_sec:.0f}s vs {ref_label} {ref_str}  →  {pct*100:.0f}%  "
                    f"⚠  SIGNIFICANTLY BELOW [{context}]")
        }
    if pct < CAUTION_BRAND_PCT:
        return "caution", {
            **base,
            "msg": (f"{match_sec:.0f}s vs {ref_label} {ref_str}  →  {pct*100:.0f}%  "
                    f"↓  below {ref_label} [{context}]")
        }
    return "ok", {
        **base,
        "msg": f"{match_sec:.0f}s vs {ref_label} {ref_str}  →  {pct*100:.0f}%  ✓  [{context}]"
    }


def compare_summary_to_stats(
    summary: pd.DataFrame,
    stats: pd.DataFrame,
    context: str,
    ref_col: str = "median",
) -> dict[str, list]:
    """Compare match summary to a stats table. ref_col: median or mean."""
    lookup = {
        (r["Brand"], r["Asset"]): r
        for _, r in stats.iterrows()
    }

    buckets: dict[str, list] = {
        "low": [], "caution": [], "high": [], "ok": [], "no_ref": [],
    }

    seen = set()
    for _, row in summary.iterrows():
        brand, asset = row["Brand"], row["Asset"]
        match_sec = float(row["Total_Sec"])
        key = (brand, asset)
        seen.add(key)

        if key in lookup:
            st = lookup[key]
            ref = float(st[ref_col])
            n = int(st["count"])
        else:
            ref, n = None, 0

        bucket, item = _classify_row(
            brand, asset, match_sec, ref, ref_col, n, context
        )
        buckets[bucket].append(item)

    # Placements in pool but zero in this match — informational only
    for key, st in lookup.items():
        if key not in seen and float(st[ref_col]) > 0:
            brand, asset = key
            buckets.setdefault("pool_only", []).append({
                "brand": brand, "asset": asset,
                "match_sec": 0.0, "ref": float(st[ref_col]),
                "n": int(st["count"]), "context": context,
                "msg": (f"0s in this match vs {ref_col} {st[ref_col]:.0f}s "
                        f"(n={int(st['count'])}) in pool — absent here"),
            })

    for k in buckets:
        buckets[k] = _by_match_sec_desc(buckets[k])
    return buckets


def _print_bucket_section(title: str, color: str, symbol: str, items: list):
    if not items:
        return
    sep = "─" * 72
    print(f"{color}{BOLD}{symbol}  {title}  ({len(items)}){RESET}")
    print(f"{DIM}{sep}{RESET}")
    for it in items:
        print(f"  {color}{it.get('symbol', '·')} {it['brand']} — {it['asset']}{RESET}")
        print(f"     {it['msg']}")
    print()


def print_comparison_report(
    match_name: str,
    venue: str,
    all_sources: list[str],
    venue_sources: list[str],
    all_buckets: dict[str, list],
    venue_buckets: dict[str, list],
):
    sep = "═" * 72
    print(f"\n{BOLD}{sep}{RESET}")
    print(f"{BOLD}  MEDIAN CHECK — vs current results pool{RESET}")
    print(f"  Match:   {CYAN}{match_name}{RESET}")
    print(f"  Venue:   {CYAN}{venue}{RESET}")
    print(f"  Pool A:  {len(all_sources)} file(s) across all venues")
    print(f"  Pool B:  {len(venue_sources)} file(s) at {venue}")
    print(f"{BOLD}{sep}{RESET}\n")

    # ── Section A: all venues median ──
    print(f"{BOLD}  (A) ALL VENUES — median per placement{RESET}\n")
    _print_bucket_section("LOW", RED, "⛔", all_buckets.get("low", []))
    _print_bucket_section("CAUTIONS", YELLOW, "⚠", all_buckets.get("caution", []))
    _print_bucket_section("HIGH", MAGENTA, "📈", all_buckets.get("high", []))
    if all_buckets.get("ok"):
        print(f"{GREEN}{BOLD}✅  OK  ({len(all_buckets['ok'])}){RESET}")
        print(f"{DIM}{'─'*72}{RESET}")
        for o in all_buckets["ok"]:
            print(f"  {GREEN}· {o['brand']} — {o['asset']}{RESET}")
            print(f"     {o['msg']}")
        print()
    if all_buckets.get("no_ref"):
        print(f"{CYAN}{BOLD}ℹ  NEW vs all-venues pool  ({len(all_buckets['no_ref'])}){RESET}")
        print(f"{DIM}{'─'*72}{RESET}")
        for n in all_buckets["no_ref"]:
            print(f"  {CYAN}· {n['brand']} — {n['asset']}{RESET}  ({n['match_sec']:.0f}s)")
        print()

    # ── Section B: same venue (absolute sec vs venue median) ──
    print(f"{BOLD}  (B) SAME VENUE — absolute sec vs {venue} median{RESET}\n")
    _print_bucket_section("LOW", RED, "⛔", venue_buckets.get("low", []))
    _print_bucket_section("CAUTIONS", YELLOW, "⚠", venue_buckets.get("caution", []))
    _print_bucket_section("HIGH", MAGENTA, "📈", venue_buckets.get("high", []))
    if venue_buckets.get("ok"):
        print(f"{GREEN}{BOLD}✅  OK  ({len(venue_buckets['ok'])}){RESET}")
        print(f"{DIM}{'─'*72}{RESET}")
        for o in venue_buckets["ok"]:
            print(f"  {GREEN}· {o['brand']} — {o['asset']}{RESET}")
            print(f"     {o['msg']}")
        print()
    if venue_buckets.get("no_ref"):
        print(f"{CYAN}{BOLD}ℹ  NEW vs venue pool  ({len(venue_buckets['no_ref'])}){RESET}")
        print(f"{DIM}{'─'*72}{RESET}")
        for n in venue_buckets["no_ref"]:
            print(f"  {CYAN}· {n['brand']} — {n['asset']}{RESET}  ({n['match_sec']:.0f}s)")
        print()

    def _sum(b):
        return sum(len(b.get(k, [])) for k in ("low", "caution", "high", "ok", "no_ref"))

    print(f"{DIM}{'─'*72}{RESET}")
    print(f"  {BOLD}SUMMARY (A all venues):{RESET}  "
          f"{RED}⛔ {len(all_buckets.get('low', []))} low{RESET}  "
          f"{YELLOW}⚠ {len(all_buckets.get('caution', []))} cautions{RESET}  "
          f"{MAGENTA}📈 {len(all_buckets.get('high', []))} high{RESET}  "
          f"{GREEN}✅ {len(all_buckets.get('ok', []))} ok{RESET}  "
          f"{CYAN}ℹ {len(all_buckets.get('no_ref', []))} new{RESET}")
    print(f"  {BOLD}SUMMARY (B {venue}):{RESET}  "
          f"{RED}⛔ {len(venue_buckets.get('low', []))} low{RESET}  "
          f"{YELLOW}⚠ {len(venue_buckets.get('caution', []))} cautions{RESET}  "
          f"{MAGENTA}📈 {len(venue_buckets.get('high', []))} high{RESET}  "
          f"{GREEN}✅ {len(venue_buckets.get('ok', []))} ok{RESET}  "
          f"{CYAN}ℹ {len(venue_buckets.get('no_ref', []))} new{RESET}")
    print(f"{BOLD}{sep}{RESET}\n")


def _check_export_paths(first_file: str, out: str | None) -> tuple[str, str]:
    """Return (New, Same) export paths next to input, including match stem."""
    stem = Path(first_file).stem
    out_dir = Path(first_file).resolve().parent
    if out:
        base = Path(out).with_suffix("")
        return str(base.parent / f"{base.name}_Check_New.xlsx"), str(
            base.parent / f"{base.name}_Check_Same.xlsx"
        )
    return (
        str(out_dir / f"{stem}_Check_New.xlsx"),
        str(out_dir / f"{stem}_Check_Same.xlsx"),
    )


def export_single_report(
    match_name: str,
    comparison_title: str,
    buckets: dict,
    stats_df: pd.DataFrame,
    out_path: str,
):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    HDR = "1A2332"
    WHITE = "FFFFFF"
    R_BG, Y_BG, P_BG, G_BG, C_BG = "FFF0F0", "FFFBF0", "F3E8FF", "F0FFF4", "F0FBFF"
    DARK = "0D1117"
    TEAL = "00C896"

    status_map = [
        ("⛔ LOW", "low", R_BG, "C0392B"),
        ("⚠ CAUTION", "caution", Y_BG, "B7770D"),
        ("📈 HIGH", "high", P_BG, "6B21A8"),
        ("✅ OK", "ok", G_BG, "1A6B3C"),
        ("ℹ NEW", "no_ref", C_BG, "0D6E8F"),
    ]

    def fill(c):
        return PatternFill("solid", start_color=c, end_color=c)

    def thin():
        s = Side(style="thin", color="D0D7DE")
        return Border(left=s, right=s, top=s, bottom=s)

    wb = Workbook()
    ws = wb.active
    ws.title = comparison_title[:31]
    ws.merge_cells("A1:G1")
    ws["A1"] = f"{comparison_title} — {match_name}"
    ws["A1"].font = Font(name="Arial", bold=True, color=TEAL, size=12)
    ws["A1"].fill = fill(DARK)
    ws.row_dimensions[1].height = 24

    headers = ["Status", "Brand", "Asset", "This Match (sec)", "Ref (sec)", "N matches", "Notes"]
    widths = [12, 20, 38, 16, 14, 10, 52]
    for i, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(row=2, column=i, value=h)
        c.font = Font(name="Arial", bold=True, color=WHITE, size=10)
        c.fill = fill(HDR)
        c.border = thin()
        ws.column_dimensions[get_column_letter(i)].width = w

    rows_out = []
    for label, key, bg, _ in status_map:
        for item in buckets.get(key, []):
            rows_out.append((label, bg, item))

    ri = 3
    for label, bg, item in rows_out:
        ref = item.get("ref")
        vals = [
            label, item["brand"], item["asset"], item["match_sec"],
            ref if ref else "—", item.get("n", "—"), item.get("msg", ""),
        ]
        for ci, v in enumerate(vals, 1):
            c = ws.cell(row=ri, column=ci, value=v)
            c.fill = fill(bg)
            c.border = thin()
            c.alignment = Alignment(
                horizontal="center" if ci not in (3, 7) else "left",
                vertical="center", wrap_text=(ci == 7),
            )
        ri += 1

    if not stats_df.empty:
        ri += 2
        ws.cell(row=ri, column=1, value="Reference stats (pool)").font = Font(bold=True)
        ri += 1
        for ci, h in enumerate(["Brand", "Asset", "median", "mean", "count", "min", "max"], 1):
            c = ws.cell(row=ri, column=ci, value=h)
            c.font = Font(bold=True, color=WHITE)
            c.fill = fill(HDR)
        ri += 1
        for _, st in stats_df.iterrows():
            for ci, col in enumerate(["Brand", "Asset", "median", "mean", "count", "min", "max"], 1):
                ws.cell(row=ri, column=ci, value=st[col])
            ri += 1

    ws.freeze_panes = "A3"
    wb.save(out_path)
    print(f"  {GREEN}Excel report saved:{RESET} {out_path}")


def main():
    parser = argparse.ArgumentParser(
        description="Compare new match vs median from results/ pool (all venues + same venue)"
    )
    parser.add_argument("files", nargs="+", help="New match xlsx file(s) to check")
    parser.add_argument("--venue", required=True,
                        help="Venue folder name under results/ (e.g. Edgbaston)")
    parser.add_argument("--results-dir", default=str(DEFAULT_RESULTS_DIR),
                        help=f"Root results folder (default: {DEFAULT_RESULTS_DIR})")
    parser.add_argument("--export", action="store_true", help="Export Excel report")
    parser.add_argument("--out", default=None, help="Output xlsx path")
    args = parser.parse_args()

    results_dir = Path(args.results_dir).resolve()
    exclude = set(_resolve_paths(args.files))

    print(f"\n{BOLD}i-sieve Median Checker{RESET}")
    print(f"{DIM}Loading match to check...{RESET}")
    df = load_match_files(*args.files)
    summary = summarise_match(df)

    match_name = Path(args.files[0]).stem
    if len(args.files) > 1:
        match_name += f" (+{len(args.files)-1} parts)"

    print(f"\n{DIM}Building reference pool from {results_dir}...{RESET}")
    all_files, venue_files = discover_pool_xlsx(results_dir, args.venue, exclude)

    print(f"\n{BOLD}All-venues pool ({len(all_files)} files):{RESET}")
    all_long, all_sources = load_summaries_from_files(all_files)
    all_stats = build_median_stats(all_long)

    print(f"\n{BOLD}Venue pool — {args.venue} ({len(venue_files)} files):{RESET}")
    venue_long, venue_sources = load_summaries_from_files(venue_files)
    venue_stats = build_median_stats(venue_long)

    if all_stats.empty and venue_stats.empty:
        print(f"{YELLOW}Warning:{RESET} No reference files in pool (excluding this match).")
        print(f"  Add finalised xlsx files under {results_dir}/<venue>/ first.")
        print(f"  Comparisons will show all placements as NEW.\n")

    all_buckets = compare_summary_to_stats(summary, all_stats, "all venues", "median")
    venue_buckets = compare_summary_to_stats(summary, venue_stats, args.venue, "median")

    print_comparison_report(
        match_name, args.venue, all_sources, venue_sources,
        all_buckets, venue_buckets,
    )

    if args.export:
        out_new, out_same = _check_export_paths(args.files[0], args.out)
        export_single_report(
            match_name, "New — all venues median", all_buckets, all_stats, out_new,
        )
        export_single_report(
            match_name, f"Same — {args.venue} median", venue_buckets, venue_stats, out_same,
        )


if __name__ == "__main__":
    main()
